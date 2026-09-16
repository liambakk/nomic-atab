"""Regenerate the Word brief from the supported Markdown in brief.md."""

from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'brief.md'
OUT = ROOT / 'deliverables/Ben - Liam and Atab data brief.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
FONT = 'Helvetica'
BLACK = RGBColor(0, 0, 0)

def set_fonts(rpr):
    for tag in ('spacing', 'position', 'kern'):
        for old in list(rpr.findall(qn('w:' + tag))):
            rpr.remove(old)
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.attrib.clear()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + key), FONT)
    for tag, value in [('sz', '22'), ('szCs', '22'), ('color', '000000')]:
        for old in list(rpr.findall(qn('w:' + tag))):
            rpr.remove(old)
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), value)
        rpr.append(el)

def style_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = BLACK
    set_fonts(style.element.get_or_add_rPr())
    for el in style.element.iter(qn('w:color')):
        el.attrib.clear()
        el.set(qn('w:val'), '000000')
    ppr = style.element.get_or_add_pPr()
    for el in list(ppr):
        if el.tag in (qn('w:pBdr'), qn('w:shd'), qn('w:contextualSpacing')):
            ppr.remove(el)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(0.68), Inches(0.65)
sec.left_margin = sec.right_margin = Inches(0.78)
sec.footer_distance = Inches(0.28)
for s in doc.styles:
    if hasattr(s, 'font'):
        s.font.name = FONT
        set_fonts(s.element.get_or_add_rPr())
for name, size, bold in [('Normal',11,False),('Title',11,True),('Subtitle',11,False),('Heading 1',11,True),('Heading 2',11,False),('List Bullet',11,False),('Footer',11,False)]:
    style_font(doc.styles[name], size, bold)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.08
normal.space_after = Pt(6)
normal.widow_control = True
for name, before, after in [('Title',0,6),('Subtitle',0,5),('Heading 1',11,7),('Heading 2',10,6)]:
    pf = doc.styles[name].paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.keep_with_next = True
    pf.keep_together = True
bp = doc.styles['List Bullet'].paragraph_format
bp.left_indent, bp.first_line_indent = Inches(0.15), Inches(-0.15)
bp.space_after = Pt(5)
bp.line_spacing = 1.08
bp.keep_together = True
for node in doc.part.numbering_part.element.iter(qn('w:rPr')):
    set_fonts(node)
for node in doc.part.numbering_part.element.iter(qn('w:lvlText')):
    if node.get(qn('w:val')) in ('\uf0b7','o','\uf0a7'):
        node.set(qn('w:val'), '\u2022')

def inline(p, text):
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        if part:
            r = p.add_run(part)
            r.font.name = FONT
            set_fonts(r._r.get_or_add_rPr())
            r.bold = p.style.name in ('Title', 'Heading 1')

footer = sec.footer.paragraphs[0]
footer.style = doc.styles['Footer']
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run()
set_fonts(r._r.get_or_add_rPr())
r.font.size = Pt(11)
for kind, txt in [('begin',None),(None,'PAGE'),('end',None)]:
    el = OxmlElement('w:instrText' if kind is None else 'w:fldChar')
    if kind:
        el.set(qn('w:fldCharType'),kind)
    else:
        el.text = txt
    r._r.append(el)

TABLE_WIDTH = 6.94

def add_table(rows):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [1.48,TABLE_WIDTH-1.48]
    for col,w in zip(t.columns,widths):
        col.width = Inches(w)
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:
            b.set(qn('w:'+k),v)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for ri,values in enumerate(rows):
        row=t.add_row()
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci,text in enumerate(values):
            cell=row.cells[ci]
            cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side,value in [('top','75'),('bottom','75'),('left','110'),('right','110')]:
                el=OxmlElement('w:'+side)
                el.set(qn('w:w'),value)
                el.set(qn('w:type'),'dxa')
                margins.append(el)
            tcpr.append(margins)
            shade=OxmlElement('w:shd')
            shade.set(qn('w:fill'),'EEEEEE' if ri == 0 else ('F5F5F5' if ri % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            p=cell.paragraphs[0]
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=(ri==0)
            inline(p,text)
            for run in p.runs:
                run.font.size=Pt(11)
                run.bold=False
    spacer=doc.add_paragraph()
    spacer.paragraph_format.space_after=Pt(0)
    spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=Pt(7)
    spacer.add_run().font.size=Pt(11)

lines=SOURCE.read_text().splitlines()
i=0
pagebreak=False
while i<len(lines):
    line=lines[i]
    i+=1
    if not line.strip():
        continue
    if line == '<!-- pagebreak -->':
        pagebreak=True
        continue
    if line.startswith('|'):
        rows=[]
        while True:
            if not re.match(r'^\|[\s:|\-]+\|$',line):
                rows.append([x.strip() for x in line.strip('|').split('|')])
            if i>=len(lines) or not lines[i].startswith('|'):
                break
            line=lines[i]
            i+=1
        add_table(rows)
        continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Title'); text=line[2:]
    elif line.startswith('## '):
        p=doc.add_paragraph(style='Heading 1'); text=line[3:]
    elif line.startswith('### '):
        p=doc.add_paragraph(style='Heading 2'); text=line[4:]
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); text=line[2:]
    elif line == 'Liam and Atab':
        p=doc.add_paragraph(style='Subtitle'); text=line
    else:
        p=doc.add_paragraph(); text=line
    inline(p,text)
    if line.startswith('Prepared for'):
        p.paragraph_format.space_after=Pt(10)
        for r in p.runs:
            r.font.size=Pt(11)
    if pagebreak:
        p.paragraph_format.page_break_before=True
        p.paragraph_format.space_before=Pt(0)
        pagebreak=False

doc.core_properties.title='Business data handoff brief'
doc.core_properties.subject='Liam and Atab'
doc.core_properties.author='Liam Bakker'
doc.core_properties.last_modified_by='Liam Bakker'
doc.core_properties.comments=''
doc.save(OUT)
print(OUT)
